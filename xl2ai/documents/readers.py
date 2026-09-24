"""Direct readers: one function per format, each returning a `Doc`. No Office needed.

Libraries are optional and imported lazily (a missing one only disables its own format, with a clear message):
python-docx (.docx), python-pptx (.pptx), pdfplumber (.pdf), olefile (.msg). E-mail (.eml), text, Markdown, CSV and
HTML use the standard library only. Rights-managed files and legacy binary .doc/.ppt go through Office itself
(`com_readers.py`, Windows).
"""
from __future__ import annotations

import csv
import datetime as dt
import email
import email.policy
import io
import os
import re
import struct
from html.parser import HTMLParser

from .model import Doc

TEXT_ENCODINGS = ("utf-8-sig", "cp1256", "latin-1")        # cp1256: Arabic Windows, common in older exports


class ReaderMissing(RuntimeError):
    """The library for this format is not installed (message says which one)."""


def _need(module, package):
    try:
        return __import__(module)
    except ImportError:
        raise ReaderMissing(f"reading this format needs the '{package}' package (pip install {package})") from None


def _decode(data):
    for enc in TEXT_ENCODINGS:
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _paragraphs(text):
    return [p.strip() for p in re.split(r"\n\s*\n", text or "") if p.strip()]


# ---- Word ------------------------------------------------------------------------------------------------------
HEADING_STYLE = re.compile(r"^(heading|عنوان)\s*(\d)", re.I)


def read_docx(path, data=None):
    docx = _need("docx", "python-docx")
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph
    d = docx.Document(io.BytesIO(data) if data is not None else path)
    doc = Doc("docx")
    cp = d.core_properties
    doc.meta = {k: v for k, v in {"title": cp.title, "author": cp.author, "subject": cp.subject,
                                  "created": cp.created and cp.created.isoformat(),
                                  "modified": cp.modified and cp.modified.isoformat()}.items() if v}
    path_stack, last_text = [], ""
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = HEADING_STYLE.match(style)
            if m or style.lower() == "title":
                level = int(m.group(2)) if m else 0
                path_stack = path_stack[:max(0, level - 1)] + [text] if level else [text]
                doc.blocks.append(_block("heading", text, section=" > ".join(path_stack[:-1]), level=level))
            else:
                kind = "list_item" if "list" in style.lower() else "paragraph"
                doc.blocks.append(_block(kind, text, section=" > ".join(path_stack)))
            last_text = text
        elif tag == "tbl":
            t = DocxTable(el, d)
            rows = []
            for r in t.rows:
                cells, prev = [], None
                for c in r.cells:                      # merged cells repeat the same cell object: keep one value
                    cells.append("" if c is prev else c.text)
                    prev = c
                rows.append(cells)
            caption = last_text if last_text and len(last_text) <= 120 else ""
            doc.add_table(rows, section=" > ".join(path_stack), caption=caption)
            last_text = ""
    return doc


# ---- PowerPoint ------------------------------------------------------------------------------------------------
def read_pptx(path, data=None):
    pptx = _need("pptx", "python-pptx")
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = pptx.Presentation(io.BytesIO(data) if data is not None else path)
    doc = Doc("pptx", pages=len(prs.slides))
    cp = prs.core_properties
    doc.meta = {k: v for k, v in {"title": cp.title, "author": cp.author,
                                  "created": cp.created and cp.created.isoformat(),
                                  "modified": cp.modified and cp.modified.isoformat()}.items() if v}

    def shapes(container):
        for sh in container:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shapes(sh.shapes)
            else:
                yield sh

    for n, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        section = f"Slide {n}" + (f": {title}" if title else "")
        if title:
            doc.blocks.append(_block("slide_title", title, page=n, section=section))
        for sh in shapes(slide.shapes):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text for c in r.cells] for r in sh.table.rows]
                doc.add_table(rows, page=n, section=section, caption=title)
            elif getattr(sh, "has_chart", False) and sh.has_chart:
                _chart_table(doc, sh.chart, n, section, title)
            elif sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip()
                    if text:
                        doc.blocks.append(_block("list_item" if para.level else "paragraph", text, page=n,
                                                 section=section))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                doc.blocks.append(_block("note", notes, page=n, section=section))
    return doc


def _chart_table(doc, chart, slide_no, section, title):
    """The numbers behind a chart (they are stored in the file): categories x series, as a table."""
    try:
        plot = chart.plots[0]
        cats = [str(c) for c in plot.categories]
        series = list(plot.series)
        if not cats or not series:
            return
        rows = [["category"] + [s.name or f"series {i + 1}" for i, s in enumerate(series)]]
        for i, cat in enumerate(cats):
            rows.append([cat] + ["" if i >= len(s.values) or s.values[i] is None else repr(float(s.values[i]))
                                 for s in series])
        chart_title = chart.chart_title.text_frame.text if chart.has_title else ""
        doc.add_table(rows, page=slide_no, section=section, caption=f"chart: {chart_title or title}")
    except Exception as e:                                   # charts vary a lot; never lose the rest of the deck
        doc.warnings.append(f"slide {slide_no}: a chart's data could not be read ({type(e).__name__})")


# ---- PDF -------------------------------------------------------------------------------------------------------
def read_pdf(path, data=None):
    pdfplumber = _need("pdfplumber", "pdfplumber")
    doc = Doc("pdf")
    with pdfplumber.open(io.BytesIO(data) if data is not None else path) as pdf:
        doc.pages = len(pdf.pages)
        info = pdf.metadata or {}
        doc.meta = {k.lower(): str(v) for k, v in info.items() if k in ("Title", "Author", "Subject", "CreationDate",
                                                                         "ModDate", "Producer") and v}
        scanned = []
        for n, page in enumerate(pdf.pages, 1):
            tables = page.find_tables()
            boxes = [t.bbox for t in tables]
            for t in tables:
                doc.add_table(t.extract(), page=n)

            def outside(obj, boxes=boxes):
                x = (obj.get("x0", 0) + obj.get("x1", 0)) / 2
                y = (obj.get("top", 0) + obj.get("bottom", 0)) / 2
                return not any(b[0] <= x <= b[2] and b[1] <= y <= b[3] for b in boxes)

            text = (page.filter(outside) if boxes else page).extract_text() or ""
            if not text.strip() and not tables:
                if page.images:
                    scanned.append(n)
                continue
            for para in _pdf_paragraphs(text):
                doc.blocks.append(_block("paragraph", para, page=n))
        if scanned:
            doc.warnings.append(f"{len(scanned)} page(s) are images without a text layer (scanned): "
                                f"{_ranges(scanned)}; their text was not read (no OCR)")
    return doc


def _pdf_paragraphs(text):
    """pdf text is line-broken by layout: join lines into paragraphs at blank lines or sentence ends."""
    out, cur = [], []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            if cur:
                out.append(" ".join(cur))
                cur = []
            continue
        cur.append(line)
        if line.endswith((".", ":", "؟", "?", "!")) and len(" ".join(cur)) > 200:
            out.append(" ".join(cur))
            cur = []
    if cur:
        out.append(" ".join(cur))
    return out


def _ranges(nums):
    out, start, prev = [], None, None
    for n in nums:
        if start is None:
            start = prev = n
        elif n == prev + 1:
            prev = n
        else:
            out.append(f"{start}-{prev}" if prev != start else str(start))
            start = prev = n
    if start is not None:
        out.append(f"{start}-{prev}" if prev != start else str(start))
    return ", ".join(out)


# ---- HTML (pages and e-mail bodies) --------------------------------------------------------------------------
class _Html(HTMLParser):
    BLOCK = {"p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6", "br", "tr", "section", "article", "blockquote"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.doc = Doc("html")
        self.buf, self.tag, self.skip = [], None, 0
        self.table, self.row, self.cell, self.in_cell = None, None, [], False
        self.path = []

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "head"):
            self.skip += 1
        elif tag == "table":
            self.flush()
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.row = []
        elif tag in ("td", "th") and self.row is not None:
            self.in_cell, self.cell = True, []
        elif tag in self.BLOCK:
            self.flush()
            self.tag = tag

    def handle_endtag(self, tag):
        if tag in ("script", "style", "head"):
            self.skip = max(0, self.skip - 1)
        elif tag in ("td", "th") and self.in_cell:
            self.row.append(" ".join("".join(self.cell).split()))
            self.in_cell = False
        elif tag == "tr" and self.row is not None and self.table is not None:
            self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table is not None:
            self.doc.add_table(self.table, section=" > ".join(self.path))
            self.table = None
        elif tag in self.BLOCK:
            self.flush()

    def handle_data(self, data):
        if self.skip:
            return
        if self.in_cell:
            self.cell.append(data)
        elif self.table is None:
            self.buf.append(data)

    def flush(self):
        text = " ".join("".join(self.buf).split())
        self.buf = []
        if not text:
            return
        if self.tag and self.tag[0] == "h" and self.tag[1:].isdigit():
            level = int(self.tag[1:])
            self.path = self.path[:level - 1] + [text]
            self.doc.blocks.append(_block("heading", text, level=level, section=" > ".join(self.path[:-1])))
        else:
            self.doc.blocks.append(_block("list_item" if self.tag == "li" else "paragraph", text,
                                          section=" > ".join(self.path)))
        self.tag = None


def read_html(path, data=None, text=None):
    parser = _Html()
    parser.feed(text if text is not None else _decode(data if data is not None else open(path, "rb").read()))
    parser.flush()
    parser.close()
    return parser.doc


# ---- plain text, Markdown, CSV ---------------------------------------------------------------------------------
def read_text(path, data=None, kind="txt"):
    raw = _decode(data if data is not None else open(path, "rb").read())
    doc = Doc(kind)
    if kind == "csv":
        try:
            dialect = csv.Sniffer().sniff(raw[:20000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        doc.add_table(list(csv.reader(io.StringIO(raw), dialect)), caption=os.path.basename(str(path or "")))
        return doc
    path_stack, table = [], []
    for chunk in _paragraphs(raw):
        lines = chunk.splitlines()
        if kind == "md" and all(line.strip().startswith("|") for line in lines) and len(lines) >= 2:
            rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines
                    if not re.match(r"^\s*\|?\s*:?-{2,}", line)]
            doc.add_table(rows, section=" > ".join(path_stack))
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", lines[0]) if kind == "md" else None
        if m:
            level = len(m.group(1))
            path_stack = path_stack[:level - 1] + [m.group(2).strip()]
            doc.blocks.append(_block("heading", m.group(2).strip(), level=level,
                                     section=" > ".join(path_stack[:-1])))
            rest = "\n".join(lines[1:]).strip()
            if rest:
                doc.blocks.append(_block("paragraph", rest, section=" > ".join(path_stack)))
        else:
            doc.blocks.append(_block("paragraph", chunk, section=" > ".join(path_stack)))
    return doc


# ---- e-mail: .eml ----------------------------------------------------------------------------------------------
def read_eml(path, data=None):
    msg = email.message_from_bytes(data if data is not None else open(path, "rb").read(), policy=email.policy.default)
    doc = Doc("eml")
    doc.meta = {k: str(msg[h]) for k, h in (("from", "From"), ("to", "To"), ("cc", "Cc"), ("subject", "Subject"),
                                            ("date", "Date"), ("message_id", "Message-ID")) if msg[h]}
    try:
        doc.meta["sent"] = email.utils.parsedate_to_datetime(msg["Date"]).isoformat() if msg["Date"] else None
    except (TypeError, ValueError):
        pass
    _email_header_block(doc)
    body = msg.get_body(preferencelist=("plain", "html"))
    if body is not None:
        content = body.get_content()
        if body.get_content_subtype() == "html":
            _merge_body(doc, read_html(None, text=content))
        else:
            for para in _paragraphs(content):
                doc.blocks.append(_block("email_body", para))
    for part in msg.iter_attachments():
        name = part.get_filename() or "attachment"
        payload = part.get_payload(decode=True)
        if payload:
            doc.attachments.append(_attachment(name, payload))
    return doc


def _email_header_block(doc):
    m = doc.meta
    head = "\n".join(f"{k}: {m[k]}" for k in ("from", "to", "cc", "date", "subject") if m.get(k))
    if head:
        doc.blocks.append(_block("email_header", head))


def _merge_body(doc, html_doc):
    for b in html_doc.blocks:
        b.kind = "email_body" if b.kind == "paragraph" else b.kind
        doc.blocks.append(b)
    for t in html_doc.tables:
        doc.add_table(t.rows, caption="table in e-mail body")


def _attachment(name, payload):
    from .model import Attachment
    return Attachment(name=os.path.basename(name.replace("\\", "/")) or "attachment", data=payload)


# ---- e-mail: Outlook .msg (OLE compound file, read directly -- no Outlook, no fragile dependency) --------------
PT_UNICODE, PT_STRING8, PT_BINARY = "001F", "001E", "0102"
MSG_PROPS = {"subject": "0037", "body": "1000", "sender_name": "0C1A", "sender_email": "0C1F",
             "sender_smtp": "5D01", "to": "0E04", "cc": "0E03", "message_id": "1035"}


def _msg_string(ole, storage, prop):
    for typ in (PT_UNICODE, PT_STRING8):
        name = f"{storage}__substg1.0_{prop}{typ}"
        if ole.exists(name):
            raw = ole.openstream(name).read()
            return raw.decode("utf-16-le", "replace").rstrip("\x00") if typ == PT_UNICODE else _decode(raw).rstrip("\x00")
    return None


def _msg_time(ole, storage=""):
    """Sent (0x0039) or delivered (0x0E06) time from the fixed-size property stream (FILETIME)."""
    name = f"{storage}__properties_version1.0"
    if not ole.exists(name):
        return None
    raw = ole.openstream(name).read()
    header = 32 if not storage else 8
    found = {}
    for off in range(header, len(raw) - 15, 16):
        ptype, pid = struct.unpack_from("<HH", raw, off)
        if ptype == 0x0040 and pid in (0x0039, 0x0E06):
            ft = struct.unpack_from("<Q", raw, off + 8)[0]
            if ft:
                found[pid] = dt.datetime(1601, 1, 1) + dt.timedelta(microseconds=ft // 10)
    t = found.get(0x0039) or found.get(0x0E06)
    return t.isoformat() if t else None


def read_msg(path, data=None):
    olefile = _need("olefile", "olefile")
    ole = olefile.OleFileIO(io.BytesIO(data) if data is not None else path)
    try:
        doc = Doc("msg")
        m = {k: _msg_string(ole, "", pid) for k, pid in MSG_PROPS.items()}
        sender = m.get("sender_smtp") or m.get("sender_email")
        frm = f"{m['sender_name']} <{sender}>" if m.get("sender_name") and sender else (m.get("sender_name") or sender)
        doc.meta = {k: v for k, v in {"from": frm, "to": m.get("to"), "cc": m.get("cc"), "subject": m.get("subject"),
                                      "sent": _msg_time(ole), "message_id": m.get("message_id")}.items() if v}
        if doc.meta.get("sent"):
            doc.meta["date"] = doc.meta["sent"]
        _email_header_block(doc)
        body = m.get("body")
        html_name = f"__substg1.0_1013{PT_BINARY}"
        if body:
            for para in _paragraphs(body):
                doc.blocks.append(_block("email_body", para))
        elif ole.exists(html_name):
            _merge_body(doc, read_html(None, text=_decode(ole.openstream(html_name).read())))
        for entry in ole.listdir(streams=False, storages=True):
            top = entry[0]
            if len(entry) != 1 or not top.startswith("__attach_version1.0_#"):
                continue
            st = top + "/"
            name = _msg_string(ole, st, "3707") or _msg_string(ole, st, "3704") or "attachment"
            data_name = f"{st}__substg1.0_3701{PT_BINARY}"
            if ole.exists(data_name):
                doc.attachments.append(_attachment(name, ole.openstream(data_name).read()))
            else:
                doc.warnings.append(f"attachment '{name}' is an embedded item (e.g. a forwarded message) and was not read")
        if any(a.name.lower() == "message.rpmsg" for a in doc.attachments):
            doc.warnings.append("this e-mail is rights-protected (message.rpmsg): its body is encrypted; only Outlook "
                                "on Windows can read it")
        return doc
    finally:
        ole.close()


# ---- dispatch ----------------------------------------------------------------------------------------------------
READERS = {"docx": read_docx, "docm": read_docx, "pptx": read_pptx, "pptm": read_pptx, "pdf": read_pdf,
           "eml": read_eml, "msg": read_msg, "html": read_html, "htm": read_html,
           "txt": lambda p, d=None: read_text(p, d, "txt"), "md": lambda p, d=None: read_text(p, d, "md"),
           "csv": lambda p, d=None: read_text(p, d, "csv")}
DOC_KINDS = set(READERS) | {"doc", "ppt"}                # legacy binaries: Office only (com_readers)


def kind_of(name):
    return os.path.splitext(str(name))[1].lower().lstrip(".")


def read_document(path, data=None, kind=None):
    kind = kind or kind_of(path)
    reader = READERS.get(kind)
    if reader is None:
        raise ReaderMissing(f".{kind} files are read through Microsoft Office (Windows); no direct reader exists")
    doc = reader(path, data)
    doc.kind = kind if kind not in ("docm", "pptm", "htm") else doc.kind
    return doc


def _block(kind, text, page=None, section="", level=None):
    from .model import Block
    return Block(kind=kind, text=text, page=page, section=section, level=level)
