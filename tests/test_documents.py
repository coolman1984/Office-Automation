"""Documents end to end: Word, PowerPoint (with chart data), PDF, e-mail with Word and Excel attachments, Markdown,
CSV and HTML -> refresh -> searchable text, typed tables, entities linked to data rows, grounded records.
Outlook .msg and the Office (COM) readers are tested against fake objects: no Outlook/Office here."""
import datetime as dt
import email.message
import io
import json
import os
import struct
import sys
import tempfile
import types
import unittest
from unittest import mock

try:
    import docx
    import openpyxl
    import pdfplumber  # noqa: F401
    import pptx
    import python_calamine  # noqa: F401
    from reportlab.pdfgen import canvas
    HAVE_DEPS = True
except ImportError:
    HAVE_DEPS = False

from xl2ai.core.log import silence
from xl2ai.documents.entities import find_entities
from xl2ai.documents.store import parse_number


class TestPureParts(unittest.TestCase):
    def test_entities(self):
        found = {(k, n) for k, _, n, _, _ in find_entities(
            "فاتورة INV-000123 بتاريخ ١٥/٠٣/٢٠٢٤ بقيمة 4,250.50 جنيه، خصم 12.5% - ahmed@x.com 01001234567")}
        self.assertTrue({("code", "INV-000123"), ("date", "2024-03-15"), ("money", "4250.50 EGP"),
                         ("percent", "12.5%"), ("email", "ahmed@x.com"), ("phone", "01001234567")} <= found)

    def test_parse_number(self):
        self.assertEqual(parse_number("1,234.50"), 1234.5)
        self.assertEqual(parse_number("(300)"), -300)
        self.assertEqual(parse_number("12%"), 0.12)
        self.assertEqual(parse_number("٣٥٠٠"), 3500)
        self.assertIsNone(parse_number("P-12"))
        self.assertIsNone(parse_number("1.2.3"))


def _folder(root):
    f = os.path.join(root, "Mixed")
    os.makedirs(f)
    d = docx.Document()
    d.core_properties.title = "Supply contract"
    d.add_heading("Supply contract 2024", 1)
    d.add_paragraph("Agreed with customer C100 for 250,000 EGP on 15/03/2024. Invoice INV-000002.")
    d.add_heading("Prices", 2)
    d.add_paragraph("Agreed prices:")
    t = d.add_table(rows=3, cols=2)
    for r, row in enumerate([["product", "price"], ["oil", "1,085.50"], ["rice", "40"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.save(os.path.join(f, "contract.docx"))
    buf = io.BytesIO()
    d.save(buf)

    p = pptx.Presentation()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Branch results"
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    cd = CategoryChartData()
    cd.categories = ["Cairo", "Giza"]
    cd.add_series("2024", (500.0, 650.5))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(6), Inches(4), cd)
    s.notes_slide.notes_text_frame.text = "Giza grew because of order INV-000002"
    p.save(os.path.join(f, "results.pptx"))

    c = canvas.Canvas(os.path.join(f, "invoice.pdf"))
    c.drawString(72, 800, "Invoice INV-000001 dated 2024-02-10 for customer C101 total 4,250.00 EGP")
    c.showPage()
    c.save()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["invoice_no", "customer", "amount"])
    for i in range(1, 11):
        ws.append([f"INV-{i:06d}", f"C{100 + i % 3}", 100.0 * i])
    wb.save(os.path.join(f, "orders.xlsx"))
    xbuf = io.BytesIO()
    wb.save(xbuf)

    m = email.message.EmailMessage()
    m["From"] = "Ahmed <ahmed@supplier.com>"
    m["To"] = "sales@company.com"
    m["Subject"] = "Order confirmation"
    m["Date"] = "Tue, 12 Mar 2024 10:00:00 +0200"
    m.set_content("Hello,\n\nWe confirm order INV-000003 for customer C102: 30 units of oil at 85 EGP, "
                  "total 2,550 EGP, delivery on 2024-03-20.\n\nRegards")
    m.add_attachment(buf.getvalue(), maintype="application", subtype="octet-stream", filename="contract copy.docx")
    m.add_attachment(xbuf.getvalue(), maintype="application", subtype="octet-stream", filename="prices.xlsx")
    with open(os.path.join(f, "order.eml"), "wb") as fh:
        fh.write(bytes(m))
    with open(os.path.join(f, "notes.md"), "w", encoding="utf-8") as fh:
        fh.write("# ملاحظات الاجتماع\n\nتم مراجعة العميل C100 والقاهرة\n\n| branch | target |\n|---|---|\n"
                 "| Cairo | 6,000 |\n| Giza | 6,500 |\n")
    with open(os.path.join(f, "targets.csv"), "w", encoding="utf-8") as fh:
        fh.write("branch;target\nCairo;6000\nGiza;6500\n")
    with open(os.path.join(f, "page.html"), "w", encoding="utf-8") as fh:
        fh.write("<html><head><style>x{}</style></head><body><h1>Price list</h1><p>Valid from 2024-01-01</p>"
                 "<table><tr><th>item</th><th>price</th></tr><tr><td>oil</td><td>85</td></tr>"
                 "<tr><td>rice</td><td>40</td></tr></table></body></html>")
    return f


@unittest.skipUnless(HAVE_DEPS, "document libraries not installed")
class TestDocumentsEndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from xl2ai.core.config import load_config
        from xl2ai.refresh import run_refresh
        from xl2ai.workspace import open_workspace
        cls.tmp = tempfile.TemporaryDirectory()
        cls.folder = _folder(cls.tmp.name)
        cfg_path = open_workspace(cls.folder, engine="direct")["config"]
        prev = silence(True)
        try:
            cls.cfg = load_config(cfg_path)
            cls.code, cls.refresh_run = run_refresh(cls.cfg)
        finally:
            silence(prev)
        import sqlite3
        cls.cat = sqlite3.connect(os.path.join(cls.refresh_run.dir, "catalog.db"))

    @classmethod
    def tearDownClass(cls):
        cls.cat.close()
        cls.tmp.cleanup()

    def test_run_passed_with_every_format(self):
        self.assertEqual(self.code, 0, self.refresh_run.m.get("stages"))
        kinds = {k for k, in self.cat.execute("SELECT kind FROM _documents")}
        self.assertEqual(kinds, {"docx", "pptx", "pdf", "eml", "md", "csv", "html"})

    def test_tables_from_documents_are_typed_and_queryable(self):
        from xl2ai.query import sql
        out = sql(self.cfg, "*", "SELECT price FROM contract_table_1_Agreed_prices ORDER BY price")
        self.assertEqual([r[0] for r in out["rows"]], [40, 1085.5])
        names = {n for n, in self.cat.execute("SELECT sheet_name FROM _tables")}
        self.assertIn("p1 table 1 -- chart: Branch results", names)             # data behind a PowerPoint chart
        self.assertIn("prices.xlsx > Orders", names)                            # Excel attached to an e-mail
        self.assertIn("contract copy.docx table 1 -- Agreed prices:", names)    # Word attached to an e-mail

    def test_search_is_located_and_arabic_tolerant(self):
        from xl2ai.documents.tools import search
        hits = search(self.cfg, "oil units")["rows"]
        self.assertEqual(hits[0][2], "email_body")
        hits = search(self.cfg, "القاهره")["rows"]                              # ة typed as ه
        self.assertEqual(len(hits), 1)
        hits = search(self.cfg, "grew")["rows"]
        self.assertIn("Slide 1: Branch results", hits[0][3])

    def test_read_around_a_hit(self):
        from xl2ai.documents.tools import read, search
        hit = search(self.cfg, "INV-000001")["rows"][0]
        out = read(self.cfg, hit[1], block=hit[0], around=1)
        self.assertEqual(out["document"]["kind"], "pdf")
        self.assertIn("4,250.00 EGP", out["rows"][0][3])

    def test_mentions_link_documents_to_rows(self):
        rows = self.cat.execute("""SELECT m.normalized, t.sheet_name, m.column_name, m.first_row FROM _mentions m
                                   JOIN _tables t ON t.table_id = m.table_id
                                   JOIN _documents d ON d.source_id = m.source_id
                                   WHERE d.kind = 'pdf'""").fetchall()
        self.assertIn(("INV-000001", "Orders", "invoice_no", 2), rows)

    def test_email_meta(self):
        sender, subject, sent = self.cat.execute(
            "SELECT sender, subject, sent FROM _documents WHERE kind='eml'").fetchone()
        self.assertEqual((sender, subject), ("Ahmed <ahmed@supplier.com>", "Order confirmation"))
        self.assertTrue(sent.startswith("2024-03-12"))

    def test_save_records_is_grounded(self):
        from xl2ai.documents.tools import save_records, search
        from xl2ai.query import sql
        block = search(self.cfg, "confirm order", kind="email_body")["rows"][0][0]
        quote = "order INV-000003 for customer C102: 30 units of oil at 85 EGP, total 2,550 EGP"
        out = save_records(self.cfg, "email_orders", [
            {"invoice_no": "INV-000003", "customer": "C102", "qty": 30, "total": 2550,
             "_source": {"block_id": block, "quote": quote}},
            {"invoice_no": "INV-000009", "customer": "C102", "_source": {"block_id": block, "quote": quote}},
            {"invoice_no": "INV-000003", "_source": {"block_id": block, "quote": "an invented sentence"}},
            {"total": 9999, "_source": {"block_id": block, "quote": quote}},
            {"invoice_no": "INV-000003"},
        ], key="invoice_no")
        self.assertEqual(out["saved"], 1)
        reasons = [r["reason"] for r in out["rejected"]]
        self.assertTrue(any("does not appear in the quote" in r for r in reasons))      # INV-000009, 9999
        self.assertTrue(any("quote not found" in r for r in reasons))                   # invented sentence
        self.assertTrue(any("missing source" in r for r in reasons))                    # no source at all
        joined = sql(self.cfg, "*", """SELECT e.invoice_no, e.total, o.amount FROM email_orders e
                                       JOIN Orders o ON o.invoice_no = e.invoice_no""")
        self.assertEqual(joined["rows"], [["INV-000003", 2550, 300.0]])
        again = save_records(self.cfg, "email_orders", [{"invoice_no": "INV-000003", "qty": 30,
                                                          "_source": {"block_id": block, "quote": quote}}],
                             key="invoice_no")
        self.assertEqual(again["saved"], 1)
        n = sql(self.cfg, "*", "SELECT COUNT(*) FROM email_orders")["rows"][0][0]
        self.assertEqual(n, 1)                                                          # replaced by key

    def test_agent_brief_lists_documents(self):
        with open(os.path.join(self.refresh_run.dir, "ai", "agent_brief.md"), encoding="utf-8") as f:
            text = f.read()
        self.assertIn("## Documents", text)
        self.assertIn("INV-000001 is in `Orders.invoice_no`", text)


class _FakeStream:
    def __init__(self, data):
        self.data = data

    def read(self):
        return self.data


class TestMsgReader(unittest.TestCase):
    """The .msg reader against a fake OLE container (the layout Outlook writes)."""

    def test_reads_headers_body_time_and_attachments(self):
        from xl2ai.documents.readers import read_msg
        ft = int((dt.datetime(2024, 3, 12, 8, 0) - dt.datetime(1601, 1, 1)).total_seconds() * 10_000_000)
        props = b"\0" * 32 + struct.pack("<HHIQ", 0x0040, 0x0039, 0, ft)
        streams = {
            "__substg1.0_0037001F": "طلبية جديدة".encode("utf-16-le"),
            "__substg1.0_1000001F": "العميل C101 طلب 20 كرتونة.\r\n\r\nشكرا".encode("utf-16-le"),
            "__substg1.0_0C1A001F": "Ahmed".encode("utf-16-le"),
            "__substg1.0_5D01001F": "ahmed@x.com".encode("utf-16-le"),
            "__substg1.0_0E04001F": "Sales".encode("utf-16-le"),
            "__properties_version1.0": props,
            "__attach_version1.0_#00000000/__substg1.0_3707001F": "note.txt".encode("utf-16-le"),
            "__attach_version1.0_#00000000/__substg1.0_37010102": b"attached text",
        }

        class FakeOle:
            def __init__(self, _):
                pass

            def exists(self, name):
                return name in streams

            def openstream(self, name):
                return _FakeStream(streams[name])

            def listdir(self, streams=True, storages=False):
                return [["__attach_version1.0_#00000000"]]

            def close(self):
                pass

        with mock.patch.dict(sys.modules, {"olefile": types.SimpleNamespace(OleFileIO=FakeOle)}):
            doc = read_msg("x.msg")
        self.assertEqual(doc.meta["subject"], "طلبية جديدة")
        self.assertEqual(doc.meta["from"], "Ahmed <ahmed@x.com>")
        self.assertTrue(doc.meta["sent"].startswith("2024-03-12T08:00"))
        self.assertEqual([b.kind for b in doc.blocks], ["email_header", "email_body", "email_body"])
        self.assertEqual([(a.name, a.data) for a in doc.attachments], [("note.txt", b"attached text")])


class TestWordCom(unittest.TestCase):
    """read_word_com against a fake Word object model."""

    def test_headings_paragraphs_and_tables(self):
        from xl2ai.documents.com_readers import read_word_com

        class Rng:
            def __init__(self, text, in_table=False):
                self.Text, self.in_table = text, in_table

            def Information(self, what):
                return self.in_table

        class Para:
            def __init__(self, text, level, in_table=False):
                self.Range, self.OutlineLevel = Rng(text, in_table), level

        class Cell:
            def __init__(self, r, c, text):
                self.RowIndex, self.ColumnIndex, self.Range = r, c, Rng(text + "\r\x07")

        class Table:
            Rows = types.SimpleNamespace(Count=2)
            Columns = types.SimpleNamespace(Count=2)
            Range = types.SimpleNamespace(Cells=[Cell(1, 1, "item"), Cell(1, 2, "price"),
                                                 Cell(2, 1, "oil"), Cell(2, 2, "85")])

        class Document:
            Paragraphs = [Para("Contract\r", 1), Para("Terms apply.\r", 10), Para("oil\r", 10, in_table=True)]
            Tables = [Table()]

            def BuiltInDocumentProperties(self, k):
                return types.SimpleNamespace(Value="T" if k == "Title" else None)

            def Close(self, save):
                self.closed = True

        class App:
            Documents = types.SimpleNamespace(Open=lambda **kw: Document())

        doc = read_word_com("x.docx", App())
        self.assertEqual([(b.kind, b.text) for b in doc.blocks], [("heading", "Contract"), ("paragraph", "Terms apply.")])
        self.assertEqual(doc.tables[0].rows, [["item", "price"], ["oil", "85"]])


if __name__ == "__main__":
    unittest.main()
