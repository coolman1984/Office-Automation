"""Key numbers (digest), `query find`, cross-file SQL, join recipes and parallel extraction -- end to end on two real
.xlsx files (sales + a separate customers workbook), with the direct engine. No Excel needed."""
import datetime as dt
import json
import os
import sqlite3
import tempfile
import unittest

try:
    import openpyxl
    import python_calamine  # noqa: F401
    HAVE_DEPS = True
except ImportError:
    HAVE_DEPS = False

from xl2ai.core.log import silence


def _sales(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["order_id", "order_date", "customer_id", "city", "qty", "unit_price", "amount"])
    cities = ["القاهرة", "Giza", "Alex"]
    for i in range(1, 61):
        qty, price = i % 5 + 1, 10.0
        ws.append([i, dt.date(2024, 1 + (i - 1) // 20, 1 + i % 20), 100 + i % 6, cities[i % 3], qty, price,
                   qty * price])
    ws.append(["Total", None, None, None, None, None, 999999])      # a totals row that must not be counted
    # a summary report of amount by city, with one number that drifted from the data
    by_city = {}
    for i in range(1, 61):
        by_city[cities[i % 3]] = by_city.get(cities[i % 3], 0) + (i % 5 + 1) * 10.0
    rep = wb.create_sheet("City Report")
    rep.append(["city", "sales"])
    for city, value in by_city.items():
        rep.append([city, value + (5 if city == "Giza" else 0)])
    rep.append(["Total", sum(by_city.values())])
    # a long daily log with one extreme value and a missing month (for the anomaly detectors)
    log = wb.create_sheet("Daily Log")
    log.append(["day", "region", "amount"])
    for d in range(0, 240):
        day = dt.date(2023, 1, 1) + dt.timedelta(days=d)
        if day.month == 5:
            continue
        log.append([day, ["N", "S"][d % 2], 100 + (d % 7) + (50000 if d == 100 else 0)])
    wb.save(path)


def _customers(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws.append(["customer_id", "customer_name", "segment"])
    for k in range(100, 106):
        ws.append([k, f"عميل {k}", ["Retail", "Corporate"][k % 2]])
    wb.save(path)


@unittest.skipUnless(HAVE_DEPS, "python-calamine and openpyxl are needed")
class TestAgentTools(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from xl2ai.core.config import load_config
        from xl2ai.refresh import run_refresh
        cls.tmp = tempfile.TemporaryDirectory()
        root = cls.tmp.name
        _sales(os.path.join(root, "sales.xlsx"))
        _customers(os.path.join(root, "customers.xlsx"))
        with open(os.path.join(root, "xl2ai.toml"), "w") as f:
            f.write('[project]\nname = "tools"\n[extract]\nengine = "direct"\nworkers = 2\n'
                    '[[sources]]\npath = "sales.xlsx"\n[[sources]]\npath = "customers.xlsx"\n')
        previous = silence(True)
        try:
            cls.cfg = load_config(os.path.join(root, "xl2ai.toml"))
            cls.code, cls.refresh_run = run_refresh(cls.cfg)
        finally:
            silence(previous)
        cls.catalog = os.path.join(cls.refresh_run.dir, "catalog.db")

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def cat(self, sql, args=()):
        con = sqlite3.connect(self.catalog)
        try:
            return con.execute(sql, args).fetchall()
        finally:
            con.close()

    def orders_id(self):
        return self.cat("SELECT table_id FROM _tables WHERE sheet_name='Orders'")[0][0]

    # -- parallel extraction ------------------------------------------------------------------------------------
    def test_parallel_extraction_passed(self):
        self.assertEqual(self.code, 0)
        stage = next(s for s in self.refresh_run.m["stages"] if s["name"] == "extract")
        self.assertEqual(stage["details"]["workers"], 2)
        self.assertEqual([s["exit_code"] for s in stage["details"]["sources"]], [0, 0])

    # -- digest -------------------------------------------------------------------------------------------------
    def test_digest_totals_exclude_the_totals_row(self):
        tid = self.orders_id()
        status, excluded = self.cat("SELECT status, excluded_rows FROM _digest_tables WHERE table_id=?", (tid,))[0]
        self.assertEqual((status, excluded), ("computed", 1))
        totals = {(m, k): v for m, k, v in self.cat(
            "SELECT measure, key, value FROM _digest WHERE table_id=? AND section='total'", (tid,))}
        self.assertEqual(totals[("*", "rows")], 60)
        self.assertEqual(totals[("amount", "sum")], sum((i % 5 + 1) * 10.0 for i in range(1, 61)))

    def test_prices_are_averaged_never_added(self):
        keys = {k for k, in self.cat("SELECT key FROM _digest WHERE table_id=? AND measure='unit_price'",
                                     (self.orders_id(),))}
        self.assertTrue(keys)
        self.assertNotIn("sum", keys)

    def test_ids_are_never_summed(self):
        measures = {m for m, in self.cat("SELECT DISTINCT measure FROM _digest WHERE table_id=?", (self.orders_id(),))}
        self.assertNotIn("customer_id", measures)
        self.assertNotIn("order_id", measures)

    def test_digest_groups_and_months(self):
        tid = self.orders_id()
        groups = self.cat("SELECT key, share FROM _digest WHERE table_id=? AND section='by_group' AND dim='city'",
                          (tid,))
        self.assertEqual(len(groups), 3)
        self.assertAlmostEqual(sum(s for _, s in groups), 1.0, places=6)
        months = [k for k, in self.cat(
            "SELECT key FROM _digest WHERE table_id=? AND section='by_month' ORDER BY rank", (tid,))]
        self.assertEqual(months, ["2024-01", "2024-02", "2024-03"])

    def test_digest_sql_reproduces_the_number(self):
        from xl2ai.query import sql
        tid = self.orders_id()
        value, statement = self.cat("""SELECT value, sql FROM _digest WHERE table_id=? AND section='total'
                                       AND measure='amount' AND key='sum'""", (tid,))[0]
        sid = self.cat("SELECT source_id FROM _tables WHERE table_id=?", (tid,))[0][0]
        out = sql(self.cfg, sid, statement)
        self.assertEqual(out["rows"][0][0], value)

    def test_digest_in_briefs_and_query(self):
        from xl2ai.query import digest
        with open(os.path.join(self.refresh_run.dir, "ai", "agent_brief.md"), encoding="utf-8") as f:
            text = f.read()
        self.assertIn("Key numbers (60 rows, 1 totals row(s) excluded)", text)
        self.assertIn("(a per-unit value: not added up)", text)
        out = digest(self.cfg, "Orders", "by_month")
        self.assertTrue(out["ok"])
        self.assertEqual(len(out["rows"]), 3)
        self.assertTrue(out["evidence"][0]["sql"])

    # -- find ---------------------------------------------------------------------------------------------------
    def test_find_names_and_arabic_spelling_variants(self):
        from xl2ai.query import find, fold
        self.assertEqual(fold("القاهرة"), fold("القاهره"))
        self.assertEqual(fold("Customer_Name"), "customer name")
        kinds = {(r[0], r[2]) for r in find(self.cfg, "customer")["rows"]}
        self.assertIn(("column", "customer_id"), kinds)
        self.assertIn(("column", "customer_name"), kinds)
        hits = find(self.cfg, "القاهره", values=True)["rows"]            # ة written as ه
        self.assertTrue(any(r[0] == "data" and r[2] == "city" for r in hits))

    def test_find_values_reports_coverage(self):
        from xl2ai.query import find
        out = find(self.cfg, "عميل 103", values=True)
        self.assertIn("searched", out["hint"])
        self.assertTrue(any(r[0] == "data" and r[2] == "customer_name" for r in out["rows"]))

    # -- cross-file SQL and joins -------------------------------------------------------------------------------
    def test_cross_file_join(self):
        from xl2ai.query import source_alias, sql
        srcs = dict(self.cat("SELECT t.sheet_name, t.source_id FROM _tables t"))
        a, b = source_alias(srcs["Orders"]), source_alias(srcs["Customers"])
        out = sql(self.cfg, "*", f"""SELECT c.segment, SUM(o.amount) FROM {a}.Orders o
                                     JOIN {b}.Customers c ON c.customer_id = o.customer_id
                                     WHERE o.order_id IS NOT NULL GROUP BY c.segment ORDER BY c.segment""")
        self.assertTrue(out["ok"])
        self.assertEqual([r[0] for r in out["rows"]], ["Corporate", "Retail"])
        self.assertIn("aliases", out["hint"])

    def test_plain_names_leave_totals_rows_out(self):
        from xl2ai.query import source_alias, sql
        clean = sql(self.cfg, "*", "SELECT SUM(amount), COUNT(*) FROM Orders")
        self.assertEqual(clean["rows"][0], [sum((i % 5 + 1) * 10.0 for i in range(1, 61)), 60])
        self.assertIn("totals rows already left out of: Orders (-1)", clean["hint"])
        sid = self.cat("SELECT source_id FROM _tables WHERE sheet_name='Orders'")[0][0]
        raw = sql(self.cfg, "*", f"SELECT COUNT(*) FROM {source_alias(sid)}.Orders")
        self.assertEqual(raw["rows"][0][0], 61)

    def test_cross_file_sql_stays_read_only(self):
        from xl2ai.core.errors import Xl2aiError
        from xl2ai.query import sql
        for bad in ("SELECT * FROM pragma_table_info('x')", "DELETE FROM x", "SELECT 1; SELECT 2"):
            with self.assertRaises(Xl2aiError):
                sql(self.cfg, "*", bad)

    def test_join_recipe_in_agent_brief(self):
        with open(os.path.join(self.refresh_run.dir, "ai", "agent_brief.md"), encoding="utf-8") as f:
            text = f.read()
        self.assertIn("## How the tables connect", text)
        self.assertIn('`Orders.customer_id` -> `Customers.customer_id`', text)
        self.assertIn('query sql "*"', text)


    # -- reconciliation and anomalies ---------------------------------------------------------------------------
    def test_report_is_checked_against_raw_data(self):
        rows = self.cat("""SELECT r.report_column, t.sheet_name, r.source_column, r.dim_column, r.agg, r.compared,
                                  r.matched, r.status, r.mismatches
                           FROM _reconciliation r JOIN _tables t ON t.table_id=r.source_table_id
                           JOIN _tables rt ON rt.table_id=r.report_table_id WHERE rt.sheet_name='City Report'""")
        self.assertEqual(len(rows), 1)
        col, src, measure, dim, agg, compared, matched, status, bad = rows[0]
        self.assertEqual((col, src, measure, dim, agg), ("sales", "Orders", "amount", "city", "SUM"))
        self.assertEqual((compared, matched, status), (4, 3, "partial"))       # 3 cities + the Total row
        self.assertEqual([b["label"] for b in json.loads(bad)], ["Giza"])

    def test_disagreement_is_a_gap(self):
        from xl2ai.brief import build_brief
        out, code = build_brief(self.cfg)
        gaps = [g for g in out["gaps"] if g["kind"] == "report_disagrees_with_data"]
        self.assertEqual(len(gaps), 1)
        self.assertIn("Giza", gaps[0]["message"])
        with open(os.path.join(self.refresh_run.dir, "ai", "agent_brief.md"), encoding="utf-8") as f:
            self.assertIn("1 of 4 rows disagree", f.read())

    def test_anomalies_found(self):
        tid = self.cat("SELECT table_id FROM _tables WHERE sheet_name='Daily Log'")[0][0]
        found = {(c, k): (n, json.loads(e)) for c, k, n, e in self.cat(
            "SELECT column_name, kind, count, examples FROM _anomalies WHERE table_id=?", (tid,))}
        count, examples = found[("amount", "outlier")]
        self.assertEqual(count, 1)
        self.assertEqual(examples[0][1], 50000 + 100 + 100 % 7)
        self.assertEqual(found[("day", "missing_period")][1], ["2023-05"])

    def test_clean_table_has_no_outliers(self):
        tid = self.orders_id()
        kinds = {k for k, in self.cat("SELECT kind FROM _anomalies WHERE table_id=?", (tid,))}
        self.assertNotIn("outlier", kinds)                       # the 999,999 totals row is excluded first


class TestWorkerSetting(unittest.TestCase):
    def test_workers_do_not_change_the_extract_fingerprint(self):
        from xl2ai.core.config import load_config
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "xl2ai.toml")
            with open(p, "w") as f:
                f.write('[extract]\nengine = "direct"\n')
            a = load_config(p).extract_fingerprint()
            with open(p, "w") as f:
                f.write('[extract]\nengine = "direct"\nworkers = 3\n')
            self.assertEqual(a, load_config(p).extract_fingerprint())


if __name__ == "__main__":
    unittest.main()
