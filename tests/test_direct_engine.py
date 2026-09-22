"""The direct (no-Excel) extraction engine, end to end: a real .xlsx written with openpyxl -> `xl2ai refresh`
with `[extract] engine = "direct"` -> extract database, catalog, brief, agent brief and query tools.

Runs on any OS. Skipped only when python-calamine or openpyxl is not installed (both are optional extras).
"""
import datetime as dt
import json
import os
import sqlite3
import tempfile
import unittest

try:
    import openpyxl
    import python_calamine  # noqa: F401
    HAVE_DEPS = True
except ImportError:
    HAVE_DEPS = False

from xl2ai.core.log import silence as set_silent


def _write_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw Sales"
    ws.append(["Monthly sales extract"])
    ws.append([])
    ws.append(["order_id", "order_date", "customer_id", "city", "qty", "unit_price (EGP)", "at"])
    for i in range(1, 41):
        ws.append([i, dt.date(2024, 1, 1) + dt.timedelta(days=i), 100 + i % 5, ["Cairo", "Giza"][i % 2],
                   i % 7 + 1, 10.5 + i, dt.datetime(2024, 1, 1, 8, i % 60)])
    ws["E10"].value = "#N/A"
    ws["E10"].data_type = "e"
    c = wb.create_sheet("Customers")
    c.append(["customer_id", "name"])
    for k in range(100, 105):
        c.append([k, f"Cust {k}"])
    r = wb.create_sheet("Report")
    r["B1"] = "Q1"
    r.merge_cells("B1:C1")
    r["D1"] = "Q2"
    r.merge_cells("D1:E1")
    r.append(["city", "Plan", "Actual", "Plan", "Actual"])
    r.append(["Cairo", 10, "=SUMIF('Raw Sales'!D:D,A3,'Raw Sales'!E:E)", 12, 13])
    r.append(["Giza", 11, "=SUMIF('Raw Sales'!D:D,A4,'Raw Sales'!E:E)", 14, 15])
    r.append([])
    r.append([])
    r.append(["product", "stock"])
    r.append(["A", 5])
    r.append(["B", 7])
    h = wb.create_sheet("Hidden")
    h.sheet_state = "hidden"
    h.append(["k", "v"])
    h.append(["a", 1])
    wb.create_sheet("Empty")
    wb.save(path)


@unittest.skipUnless(HAVE_DEPS, "python-calamine and openpyxl are needed for the direct engine")
class TestDirectEngine(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from xl2ai.core.config import load_config
        from xl2ai.refresh import run_refresh
        cls.tmp = tempfile.TemporaryDirectory()
        root = cls.tmp.name
        _write_workbook(os.path.join(root, "sales.xlsx"))
        with open(os.path.join(root, "xl2ai.toml"), "w") as f:
            f.write('[project]\nname = "direct"\n[extract]\nengine = "direct"\n[[sources]]\npath = "sales.xlsx"\n')
        set_silent(True)
        try:
            cls.cfg = load_config(os.path.join(root, "xl2ai.toml"))
            cls.code, cls.refresh_run = run_refresh(cls.cfg)
        finally:
            set_silent(False)
        cls.run_dir = cls.refresh_run.dir
        extract_dir = os.path.join(cls.run_dir, "extract")
        cls.db = os.path.join(extract_dir, os.listdir(extract_dir)[0])

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def q(self, sql, args=()):
        con = sqlite3.connect(self.db)
        try:
            return con.execute(sql, args).fetchall()
        finally:
            con.close()

    def test_run_is_promoted_and_verified(self):
        self.assertEqual(self.code, 0)
        self.assertEqual(self.refresh_run.m["status"], "passed")
        bad = self.q("SELECT * FROM _verification WHERE ok IS NOT 1")
        self.assertEqual(bad, [])
        self.assertGreater(len(self.q("SELECT * FROM _verification")), 10)
        notes = {r[0] for r in self.q("SELECT note FROM _verification")}
        self.assertEqual(notes, {"direct engine: independent XML cell scan"})
        meta = dict(self.q("SELECT key, value FROM _meta"))
        self.assertEqual(meta["engine"], "direct")

    def test_sheets_and_layout(self):
        log = {r[0]: r[1:] for r in self.q(
            "SELECT sheet_name, status, header_row, first_row, data_rows, visibility FROM _extraction_log")}
        self.assertEqual(log["Raw Sales"], ("extracted", 3, 1, 40, "visible"))
        self.assertEqual(log["Hidden"][0], "extracted")
        self.assertEqual(log["Hidden"][4], "hidden")
        self.assertEqual(log["Empty"][0], "skipped")
        pre = self.q("SELECT xl_row, xl_col, value FROM _sheet_preamble WHERE table_name='Raw_Sales'")
        self.assertEqual(pre, [(1, 1, "Monthly sales extract")])

    def test_types_dates_and_errors(self):
        cols = {r[0]: r[1:] for r in self.q(
            "SELECT sql_name, sql_type, kind, date_format FROM _columns WHERE table_name='Raw_Sales'")}
        self.assertEqual(cols["order_id"][:2], ("INTEGER", "int"))
        self.assertEqual(cols["order_date"], ("TEXT", "date", "date"))
        self.assertEqual(cols["at"], ("TEXT", "date", "datetime"))
        self.assertEqual(cols["unit_price_EGP"][:2], ("REAL", "real"))
        row = self.q('SELECT order_date, "at", qty FROM Raw_Sales WHERE _xl_row=4')[0]
        self.assertEqual(row, ("2024-01-02", "2024-01-01 08:01:00", 2))
        self.assertEqual(self.q("SELECT xl_row, xl_col, error FROM _cell_errors"), [(10, 5, "#N/A")])
        self.assertIsNone(self.q("SELECT qty FROM Raw_Sales WHERE _xl_row=10")[0][0])

    def test_structure_recorded(self):
        regions = self.q("SELECT region_no, first_row, last_row, header_row, kind FROM _regions WHERE table_name='Report'")
        self.assertEqual(regions, [(1, 1, 4, 2, "table"), (2, 7, 9, 7, "table")])
        groups = dict(self.q("SELECT sql_name, path FROM _header_groups WHERE table_name='Report'"))
        self.assertEqual(json.loads(groups["Plan"]), ["Q1"])
        self.assertEqual(json.loads(groups["Actual_2"]), ["Q2"])

    def test_lineage_in_catalog(self):
        con = sqlite3.connect(os.path.join(self.run_dir, "catalog.db"))
        try:
            rows = con.execute("""SELECT t.sheet_name, l.ref_sheet, tt.sheet_name, l.status, l.cells, l.method
                                  FROM _lineage l JOIN _tables t ON t.table_id=l.table_id
                                  LEFT JOIN _tables tt ON tt.table_id=l.target_table_id""").fetchall()
        finally:
            con.close()
        self.assertEqual(rows, [("Report", "Raw Sales", "Raw Sales", "resolved", 2, "all_formulas")])

    def test_brief_and_agent_brief_explain_structure(self):
        from xl2ai.brief import build_brief
        out, code = build_brief(self.cfg)
        report = next(t for t in out["tables"] if t["sheet"] == "Report")
        self.assertEqual(report["regions"], 2)
        self.assertEqual(report["readiness"], "needs_review")
        self.assertEqual([f["sheet"] for f in report["feeds_from"]], ["Raw Sales"])
        self.assertIn("several_tables_in_sheet", {g["kind"] for g in out["gaps"]})
        with open(os.path.join(self.run_dir, "ai", "agent_brief.md"), encoding="utf-8") as f:
            text = f.read()
        self.assertIn("Computed from: `Raw_Sales`", text)
        self.assertIn("Holds 2 separate tables", text)
        self.assertIn("-- under Q1", text)
        with open(os.path.join(self.run_dir, "ai", "context_pack.md"), encoding="utf-8") as f:
            self.assertIn("Computed from (formula lineage)", f.read())

    def test_query_region_reads_one_table_back_out(self):
        from xl2ai.query import meta, region
        out = region(self.cfg, "Report", 2)
        self.assertTrue(out["ok"])
        self.assertEqual(out["columns"], ["_xl_row", "product", "stock"])
        self.assertEqual(out["rows"], [[8, "A", 5.0], [9, "B", 7.0]])
        first = region(self.cfg, "Report", 1)
        self.assertEqual(first["columns"][:3], ["_xl_row", "city", "Plan"])
        self.assertEqual(first["total_rows"], 2)
        self.assertTrue(meta(self.cfg, "lineage")["rows"])
        self.assertEqual(len(meta(self.cfg, "regions")["rows"]), 2)

    def test_unchanged_source_is_reused(self):
        from xl2ai.refresh import run_refresh
        set_silent(True)
        try:
            code, run = run_refresh(self.cfg)
        finally:
            set_silent(False)
        self.assertEqual(code, 0)
        stage = next(s for s in run.m["stages"] if s["name"] == "extract")
        self.assertTrue(stage["details"]["sources"][0]["reused"])


@unittest.skipUnless(HAVE_DEPS, "python-calamine and openpyxl are needed for the direct engine")
class TestXmlScanners(unittest.TestCase):
    """The fast (whole-chunk regex) and precise (per-cell) XML scanners must agree on every count and sum."""

    def test_fast_and_precise_agree(self):
        from xl2ai.extract.direct import SheetScan, XlsxPackage
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "s.xlsx")
            _write_workbook(path)
            pkg = XlsxPackage(path)
            shared = pkg._shared()
            for sheet in ("Raw Sales", "Report"):
                with pkg.zf.open(pkg.sheet_parts[sheet]) as f:
                    xml = f.read()
                fast, precise = SheetScan(), SheetScan()
                pkg._scan_fast(xml, fast, {"row_no": 0, "shared_refs": {}, "cols": {}})
                pkg._scan_precise(xml, precise, {"row_no": 0, "shared_refs": {}, "cols": {}}, shared, 0)
                self.assertEqual(fast.count, precise.count, sheet)
                self.assertEqual(fast.total, precise.total, sheet)
                self.assertEqual(fast.errors, precise.errors, sheet)
                self.assertEqual({k: round(v, 6) for k, v in fast.sums.items()},
                                 {k: round(v, 6) for k, v in precise.sums.items()}, sheet)
                self.assertEqual(fast.formulas, precise.formulas, sheet)
            self.assertTrue(fast.formulas)                    # the Report sheet really has formulas


@unittest.skipUnless(HAVE_DEPS, "python-calamine and openpyxl are needed for the direct engine")
class TestDirectEngineRefusals(unittest.TestCase):
    def _opts(self):
        from types import SimpleNamespace
        return SimpleNamespace(verify=True, strict=False, block_cells=500_000, cache_cells=0, sheets=set(),
                               wrapper_prefixes=(b"<## DRM",), engine="direct")

    def test_drm_file_is_refused_with_a_clear_reason(self):
        from xl2ai.extract.direct import process_file_direct
        with tempfile.TemporaryDirectory() as d:
            src = os.path.join(d, "x.xlsx")
            with open(src, "wb") as f:
                f.write(b"<## DRM wrapped")
            code, results, msg = process_file_direct(src, os.path.join(d, "x.db"), self._opts())
            self.assertEqual(code, 1)
            self.assertIn("Excel", msg)
            self.assertFalse(os.path.exists(os.path.join(d, "x.db")))

    def test_corrupt_file_fails_without_leaving_a_database(self):
        from xl2ai.extract.direct import process_file_direct
        with tempfile.TemporaryDirectory() as d:
            src = os.path.join(d, "x.xlsx")
            with open(src, "wb") as f:
                f.write(b"PK\x03\x04 not really a zip")
            set_silent(True)
            try:
                code, _, msg = process_file_direct(src, os.path.join(d, "x.db"), self._opts())
            finally:
                set_silent(False)
            self.assertEqual(code, 1)
            self.assertTrue(msg)
            self.assertEqual(sorted(os.listdir(d)), ["x.xlsx"])


class TestEngineSetting(unittest.TestCase):
    def test_invalid_engine_is_a_config_error(self):
        from xl2ai.core.config import load_config
        from xl2ai.core.errors import Xl2aiError
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "xl2ai.toml")
            with open(p, "w") as f:
                f.write('[extract]\nengine = "fast"\n')
            with self.assertRaises(Xl2aiError):
                load_config(p)

    def test_auto_resolves_by_platform(self):
        from types import SimpleNamespace
        from xl2ai.extract.common import PYWIN32_AVAILABLE
        from xl2ai.extract.pipeline import resolve_engine
        expected = "excel" if os.name == "nt" and PYWIN32_AVAILABLE else "direct"
        self.assertEqual(resolve_engine(SimpleNamespace(engine="auto")), expected)
        self.assertEqual(resolve_engine(SimpleNamespace(engine="direct")), "direct")


if __name__ == "__main__":
    unittest.main()
